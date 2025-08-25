#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Extract content from PPTX file and save to output.txt"""

from pptx import Presentation
import os

def extract_pptx_content(file_path):
    """Extract all content from PPTX file."""
    try:
        prs = Presentation(file_path)
        content = []
        
        content.append("=" * 60)
        content.append("SLIDEFORGE PPTX CONTENT EXTRACTION")
        content.append("=" * 60)
        content.append(f"Total slides: {len(prs.slides)}")
        content.append("")
        
        for i, slide in enumerate(prs.slides):
            content.append(f"SLIDE {i+1}:")
            content.append("-" * 30)
            
            # Extract title
            title = "No title"
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame') and shape.text_frame:
                    if shape.text_frame.paragraphs:
                        title = shape.text_frame.paragraphs[0].text.strip()
                        break
            
            content.append(f"Title: {title}")
            content.append("")
            
            # Extract all text content
            slide_content = []
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame') and shape.text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text and text != title:
                            slide_content.append(text)
            
            if slide_content:
                content.append("Content:")
                for text in slide_content:
                    content.append(f"  • {text}")
            else:
                content.append("Content: No additional content")
            
            content.append("")
            content.append("")
        
        return content
        
    except Exception as e:
        return [f"Error extracting content: {e}"]

def save_to_file(content, output_file):
    """Save content to text file."""
    try:
        with open(output_file, 'w', encoding='utf-8') as f:
            for line in content:
                f.write(line + '\n')
        print(f"Content saved to {output_file}")
    except Exception as e:
        print(f"Error saving file: {e}")

if __name__ == "__main__":
    pptx_file = "../output/research_paper_presentation (10).pptx"
    output_file = "output.txt"
    
    if os.path.exists(pptx_file):
        content = extract_pptx_content(pptx_file)
        save_to_file(content, output_file)
    else:
        print(f"PPTX file not found: {pptx_file}")
