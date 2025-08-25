#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation

# Load the PPTX file
prs = Presentation('../output/research_paper_presentation (10).pptx')

# Create output content
output_lines = []
output_lines.append("=" * 60)
output_lines.append("SLIDEFORGE PPTX CONTENT EXTRACTION")
output_lines.append("=" * 60)
output_lines.append(f"Total slides: {len(prs.slides)}")
output_lines.append("")

# Extract content from each slide
for i, slide in enumerate(prs.slides):
    output_lines.append(f"SLIDE {i+1}:")
    output_lines.append("-" * 30)
    
    # Get slide title
    title = "No title"
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            if shape.text_frame.paragraphs:
                title = shape.text_frame.paragraphs[0].text.strip()
                break
    
    output_lines.append(f"Title: {title}")
    output_lines.append("")
    
    # Get all text content
    slide_content = []
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text and text != title:
                    slide_content.append(text)
    
    if slide_content:
        output_lines.append("Content:")
        for text in slide_content:
            output_lines.append(f"  • {text}")
    else:
        output_lines.append("Content: No additional content")
    
    output_lines.append("")
    output_lines.append("")

# Save to output.txt
with open('output.txt', 'w', encoding='utf-8') as f:
    for line in output_lines:
        f.write(line + '\n')

print("Content extracted and saved to output.txt")
