from pptx import Presentation

prs = Presentation('../output/research_paper_presentation (10).pptx')
print(f'Total slides: {len(prs.slides)}')

for i, slide in enumerate(prs.slides):
    title = "No title"
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            if shape.text_frame.paragraphs:
                title = shape.text_frame.paragraphs[0].text.strip()
                break
    
    print(f'Slide {i+1}: {title}')
    
    # Check for issues
    issues = []
    if "Other —" in title or "Additional Information —" in title:
        issues.append("Bad section header")
    if "**" in title or "*" in title:
        issues.append("Raw markdown")
    if " — Key terms:" in title:
        issues.append("Inline key terms")
    
    if issues:
        print(f"  Issues: {', '.join(issues)}")
