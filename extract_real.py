#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation

# Load the ACTUAL PPTX file
prs = Presentation('../output/research_paper_presentation (11).pptx')

print("=" * 60)
print("REAL PPTX CONTENT EXTRACTION")
print("=" * 60)
print(f"Total slides: {len(prs.slides)}")
print("")

# Extract REAL content from each slide
for i, slide in enumerate(prs.slides):
    print(f"SLIDE {i+1}:")
    print("-" * 30)
    
    # Get slide title
    title = "No title"
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            if shape.text_frame.paragraphs:
                title = shape.text_frame.paragraphs[0].text.strip()
                break
    
    print(f"Title: {title}")
    print("")
    
    # Get ALL text content from the slide
    slide_content = []
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text and text != title:
                    slide_content.append(text)
    
    if slide_content:
        print("Content:")
        for text in slide_content:
            print(f"  • {text}")
    else:
        print("Content: No additional content")
    
    print("")
    print("")

print("REAL content extracted!")
