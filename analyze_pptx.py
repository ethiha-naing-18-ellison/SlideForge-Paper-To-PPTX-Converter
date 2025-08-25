#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Analyze PPTX file to identify formatting issues."""

from pptx import Presentation

prs = Presentation('../output/research_paper_presentation (10).pptx')
print(f'Total slides: {len(prs.slides)}')

issues = []

for i, slide in enumerate(prs.slides):
    title = "No title"
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            if shape.text_frame.paragraphs:
                title = shape.text_frame.paragraphs[0].text.strip()
                break
    
    print(f'Slide {i+1}: {title}')
    
    # Check for issues
    if "Other —" in title or "Additional Information —" in title:
        issues.append(f"Slide {i+1}: Bad section header '{title}'")
    if "**" in title or "*" in title:
        issues.append(f"Slide {i+1}: Raw markdown in title '{title}'")
    if " — Key terms:" in title:
        issues.append(f"Slide {i+1}: Inline key terms in title")

print("\n" + "=" * 50)
print("🔍 ISSUES FOUND:")
if issues:
    for issue in issues:
        print(f"  ❌ {issue}")
else:
    print("  ✅ No obvious issues found!")
