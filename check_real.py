#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pptx import Presentation

prs = Presentation('../output/research_paper_presentation (11).pptx')

print("REAL PPTX CONTENT:")
print("=" * 50)

for i, slide in enumerate(prs.slides):
    title = "No title"
    for shape in slide.shapes:
        if hasattr(shape, 'text_frame') and shape.text_frame:
            if shape.text_frame.paragraphs:
                title = shape.text_frame.paragraphs[0].text.strip()
                break
    
    print(f"Slide {i+1}: {title}")

print("=" * 50)
