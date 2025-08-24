#!/usr/bin/env python3
"""
Verification script to check that generated PowerPoint files don't contain "Click to add" placeholders.
"""

import sys
from pathlib import Path
from pptx import Presentation

def check_for_click_to_add(pptx_path: str) -> bool:
    """Check if a PowerPoint file contains any 'Click to add' placeholders."""
    try:
        prs = Presentation(pptx_path)
        has_placeholders = False
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip().lower()
                        if text.startswith("click to add"):
                            print(f"❌ Found 'Click to add' text: '{paragraph.text}'")
                            has_placeholders = True
        
        return has_placeholders
        
    except Exception as e:
        print(f"❌ Error reading PowerPoint file: {e}")
        return True  # Assume problematic if we can't read it

def main():
    """Check the generated PowerPoint files for placeholders."""
    print("🔍 Verifying Placeholder Cleanup in SlideForge Output")
    print("=" * 60)
    
    # Check the test output files
    test_files = [
        "test_output.pptx",
        "pagination_demo.pptx"
    ]
    
    all_clean = True
    
    for file_path in test_files:
        if Path(file_path).exists():
            print(f"\n📄 Checking: {file_path}")
            has_placeholders = check_for_click_to_add(file_path)
            
            if has_placeholders:
                print(f"❌ {file_path} contains 'Click to add' placeholders")
                all_clean = False
            else:
                print(f"✅ {file_path} is clean - no placeholders found")
        else:
            print(f"⚠️  {file_path} not found - skipping")
    
    print("\n" + "=" * 60)
    if all_clean:
        print("🎉 All PowerPoint files are clean!")
        print("✅ No 'Click to add' placeholders found")
        return 0
    else:
        print("❌ Some files still contain placeholders")
        return 1

if __name__ == "__main__":
    sys.exit(main())
