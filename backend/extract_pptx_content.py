#!/usr/bin/env python3

import sys
from pathlib import Path

# Add backend to path
backend_path = Path(__file__).parent
sys.path.insert(0, str(backend_path))

def extract_pptx_content(pptx_path: str):
    """Extract content from PPTX file."""
    try:
        from pptx import Presentation
        
        prs = Presentation(pptx_path)
        print(f"Total slides: {len(prs.slides)}")
        print("=" * 50)
        
        for i, slide in enumerate(prs.slides, 1):
            print(f"\nSLIDE {i}:")
            print("-" * 30)
            
            # Extract title
            title = ""
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
            print(f"Title: {title}")
            
            # Extract content
            content = []
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame') and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text and text != title:
                            content.append(text)
            
            if content:
                print("Content:")
                for item in content:
                    print(f"  {item}")
            else:
                print("Content: (empty)")
                
    except Exception as e:
        print(f"Error extracting content: {e}")

if __name__ == "__main__":
    pptx_path = "../test_fixed5.pptx"
    extract_pptx_content(pptx_path)
