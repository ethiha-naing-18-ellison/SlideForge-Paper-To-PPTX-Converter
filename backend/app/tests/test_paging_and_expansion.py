# backend/app/tests/test_paging_and_expansion.py
from app.services.ppt.pager import chunk_bullets
from app.services.ppt.expander import expand_to_target_slides_supplement
from pptx import Presentation

def test_chunk_bullets_respects_limits():
    bullets = [f"Point {i}" for i in range(13)]
    pages = chunk_bullets(bullets, bullets_per_slide=4, min_slides=2, max_slides=5)
    assert 2 <= len(pages) <= 5
    assert all(1 <= len(pg) for pg in pages)

def test_chunk_bullets_handles_empty_input():
    pages = chunk_bullets([], bullets_per_slide=4, min_slides=1, max_slides=3)
    assert len(pages) >= 1
    assert all(len(pg) >= 1 for pg in pages)

def test_chunk_bullets_enforces_max():
    bullets = [f"Point {i}" for i in range(20)]
    pages = chunk_bullets(bullets, bullets_per_slide=3, min_slides=1, max_slides=3)
    assert len(pages) <= 3

def test_chunk_bullets_enforces_min():
    bullets = ["Point 1", "Point 2"]
    pages = chunk_bullets(bullets, bullets_per_slide=4, min_slides=3, max_slides=5)
    assert len(pages) >= 3

def test_expand_to_target_adds_slides():
    prs = Presentation()
    # add one minimal slide
    prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    initial = len(prs.slides)
    expand_to_target_slides_supplement(prs, target_count=5, raw_text="Intro\nMethods\nResults")
    assert len(prs.slides) >= 5 and len(prs.slides) > initial

def test_expand_to_target_respects_target():
    prs = Presentation()
    # add slides to exceed target
    for _ in range(10):
        prs.slides.add_slide(prs.slide_layouts[0])
    initial = len(prs.slides)
    expand_to_target_slides_supplement(prs, target_count=5, raw_text="Test content")
    assert len(prs.slides) == initial  # Should not add more slides

def test_expand_to_target_without_raw_text():
    prs = Presentation()
    # add one minimal slide
    prs.slides.add_slide(prs.slide_layouts[0])
    initial = len(prs.slides)
    expand_to_target_slides_supplement(prs, target_count=5, raw_text=None)
    # With capped supplementary slides, we expect at least 2 more slides (Key Takeaways + Further Reading)
    assert len(prs.slides) >= 3 and len(prs.slides) > initial
