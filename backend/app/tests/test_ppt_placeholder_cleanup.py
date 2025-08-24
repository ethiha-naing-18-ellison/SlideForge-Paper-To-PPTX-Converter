import os
import pytest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

# Add the app directory to the path for imports
import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from app.services.ppt.cleanup import remove_unused_placeholders
from app.services.ppt.builder import add_blank_slide, add_title_textbox, add_titled_slide_placeholder_free
from app.models.schema import LayoutConfig

def _has_click_to_add(slide) -> bool:
    """Check if slide contains any 'Click to add' text."""
    for shp in slide.shapes:
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            txt = "".join(p.text or "" for p in shp.text_frame.paragraphs).strip().lower()
            if txt.startswith("click to add"):
                return True
    return False

def test_remove_placeholders_blank_slide():
    """Test that blank slides are already placeholder-free."""
    prs = Presentation()
    slide = add_blank_slide(prs)  # should already be placeholder-free
    assert not _has_click_to_add(slide)

def test_cleanup_removes_default_text():
    """Test that cleanup removes default placeholder text."""
    prs = Presentation()
    # Use a title+content layout intentionally to simulate legacy slides
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    s = prs.slides.add_slide(layout)
    
    # Count placeholders before cleanup
    placeholders_before = len([shp for shp in s.shapes if getattr(shp, "is_placeholder", False)])
    
    removed = remove_unused_placeholders(s)
    
    # Count placeholders after cleanup
    placeholders_after = len([shp for shp in s.shapes if getattr(shp, "is_placeholder", False)])
    
    # Should have removed some placeholders (if any existed)
    if placeholders_before > 0:
        assert removed >= 1  # Should have removed at least one placeholder
        assert placeholders_after < placeholders_before  # Should have fewer placeholders
    
    # Should not have any "Click to add" text after cleanup
    assert not _has_click_to_add(s)  # Should be clean after cleanup

def test_placeholder_free_slide_creation():
    """Test that placeholder-free slide creation works correctly."""
    prs = Presentation()
    layout_config = LayoutConfig()
    
    slide = add_titled_slide_placeholder_free(prs, layout_config, "Test Title")
    
    # Should not have any "Click to add" text
    assert not _has_click_to_add(slide)
    
    # Should have our custom title textbox
    has_title = False
    for shp in slide.shapes:
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            txt = "".join(p.text or "" for p in shp.text_frame.paragraphs).strip()
            if txt == "Test Title":
                has_title = True
                break
    
    assert has_title, "Custom title textbox should be present"

def test_add_title_textbox():
    """Test custom title textbox creation."""
    prs = Presentation()
    slide = add_blank_slide(prs)
    
    # Add a custom title textbox
    tb = add_title_textbox(slide, "Custom Title", 1.0, 1.0, 5.0, 1.0, pt=24)
    
    # Check that the textbox was created with correct text
    assert tb.has_text_frame
    assert tb.text_frame.paragraphs[0].text == "Custom Title"
    
    # Should not have any "Click to add" text
    assert not _has_click_to_add(slide)

def test_cleanup_preserves_custom_content():
    """Test that cleanup doesn't remove custom content."""
    prs = Presentation()
    slide = add_blank_slide(prs)
    
    # Add custom content
    tb = add_title_textbox(slide, "Important Content", 1.0, 1.0, 5.0, 1.0)
    
    # Run cleanup
    removed = remove_unused_placeholders(slide)
    
    # Custom content should still be there
    assert tb.text_frame.paragraphs[0].text == "Important Content"
    
    # Should not have any "Click to add" text
    assert not _has_click_to_add(slide)

def test_cleanup_handles_empty_slide():
    """Test that cleanup handles empty slides gracefully."""
    prs = Presentation()
    slide = add_blank_slide(prs)
    
    # Run cleanup on already clean slide
    removed = remove_unused_placeholders(slide)
    
    # Should not have any "Click to add" text
    assert not _has_click_to_add(slide)
    assert removed == 0  # No placeholders to remove

def test_cleanup_handles_mixed_content():
    """Test cleanup with mixed placeholder and custom content."""
    prs = Presentation()
    # Start with a layout that has placeholders
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    
    # Count placeholders before adding custom content
    placeholders_before = len([shp for shp in slide.shapes if getattr(shp, "is_placeholder", False)])
    
    # Add custom content
    tb = add_title_textbox(slide, "Custom Title", 1.0, 1.0, 5.0, 1.0)
    
    # Run cleanup
    removed = remove_unused_placeholders(slide)
    
    # Custom content should be preserved
    assert tb.text_frame.paragraphs[0].text == "Custom Title"
    
    # Should not have any "Click to add" text
    assert not _has_click_to_add(slide)
    
    # Should have removed some placeholders if they existed
    if placeholders_before > 0:
        assert removed > 0  # Should have removed some placeholders
