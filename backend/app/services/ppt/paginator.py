# backend/app/services/ppt/paginator.py
from __future__ import annotations
from typing import List, Tuple
from textwrap import wrap
from pptx.util import Inches, Pt
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT
try:
    from pptx.enum.text import MSO_AUTO_SIZE  # not always available
    HAS_AUTOFIT = True
except Exception:
    HAS_AUTOFIT = False

EMU_PER_INCH = 914400

def _emu(inches: float) -> int:
    return int(inches * EMU_PER_INCH)

def estimate_lines(text: str, box_width_in: float, font_pt: int) -> int:
    """
    Very good approximation for wrapped lines without rendering.
    chars_per_line ≈ (box_width_in * 144) / font_pt
    144 comes from 72pt/in * 2 (avg glyph ~0.5em). Tuned for common fonts.
    """
    if not text:
        return 1
    approx_cpl = max(18, int((box_width_in * 144) / max(10, font_pt)))
    wrapped = wrap(text, width=approx_cpl, break_long_words=False, replace_whitespace=False)
    return max(1, len(wrapped))

def paragraph_height_in(num_lines: int, font_pt: int, line_spacing: float) -> float:
    # height per line ~ font_pt * (line_spacing) points; convert to inches
    pts = num_lines * font_pt * line_spacing
    return pts / 72.0

def total_bullets_height_in(
    bullets: List[str],
    box_width_in: float,
    font_pt: int,
    line_spacing: float,
    para_spacing_pts: int = 6,
) -> float:
    h = 0.0
    for b in bullets:
        lines = estimate_lines(b, box_width_in, font_pt)
        h += paragraph_height_in(lines, font_pt, line_spacing)
        h += para_spacing_pts / 72.0
    return h

def split_bullets_to_fit(
    bullets: List[str],
    box_width_in: float,
    max_height_in: float,
    font_max_pt: int,
    font_min_pt: int,
    line_spacing: float,
    max_lines_per_item: int,
) -> List[Tuple[List[str], int]]:
    """
    Returns a list of (bullets_for_slide, chosen_font_pt).
    Strategy:
      1) Try to fit all bullets at decreasing font sizes (max -> min).
      2) If still doesn't fit, greedily pack bullets into multiple slides.
      3) If a single bullet is too tall (exceeds per-slide), split that bullet into continuation bullets (↳ ...).
    """
    # 1) Try global shrink at best quality
    for size in range(font_max_pt, font_min_pt - 1, -1):
        h = total_bullets_height_in(bullets, box_width_in, size, line_spacing)
        if h <= max_height_in:
            return [(bullets, size)]

    # 2) Greedy paginate
    slides: List[Tuple[List[str], int]] = []
    work = list(bullets)

    while work:
        size = font_min_pt  # we already know max won't fit; use min for packing
        current: List[str] = []
        current_h = 0.0
        for idx, b in enumerate(work):
            # Split bullet if it will exceed max_lines_per_item in this width
            lines = estimate_lines(b, box_width_in, size)
            if lines > max_lines_per_item:
                # split into chunks preserving words
                approx_cpl = max(18, int((box_width_in * 144) / max(10, size)))
                wrapped = wrap(b, width=approx_cpl, break_long_words=False, replace_whitespace=False)
                chunks = []
                while wrapped:
                    chunk_lines = wrapped[:max_lines_per_item]
                    wrapped = wrapped[max_lines_per_item:]
                    prefix = "" if not chunks else "↳ "
                    chunks.append(prefix + " ".join(chunk_lines))
                # Replace the current bullet with first chunk and put rest back
                work = chunks + work[idx + 1:]
                b = work[0]
                lines = estimate_lines(b, box_width_in, size)

            para_h = paragraph_height_in(lines, size, line_spacing) + (6 / 72.0)
            if current_h + para_h <= max_height_in:
                current.append(b)
                current_h += para_h
            else:
                break

        if not current:
            # Even a single (possibly huge) bullet doesn't fit, force-split again more aggressively
            b = work[0]
            approx_cpl = max(10, int((box_width_in * 120) / max(8, size)))
            wrapped = wrap(b, width=approx_cpl, break_long_words=True)
            half = max(1, len(wrapped)//2)
            first = " ".join(wrapped[:half])
            second = "↳ " + " ".join(wrapped[half:])
            current = [first]
            work = [second] + work[1:]
        else:
            work = work[len(current):]

        slides.append((current, size))

    return slides

def render_bullets_block(
    slide,
    left_in: float, top_in: float, width_in: float, height_in: float,
    title_text: str | None,
    bullets: List[str],
    font_pt: int,
    line_spacing: float
):
    """Render title (optional) and bullets into a text box within bounds."""
    from pptx.util import Inches, Pt

    textbox = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tf = textbox.text_frame
    tf.word_wrap = True
    if HAS_AUTOFIT:
        try:
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception:
            pass

    # Optionally prepend a title inside the same box if needed (usually title has its own box)
    if title_text:
        p0 = tf.paragraphs[0]
        p0.text = title_text
        p0.font.size = Pt(font_pt + 4)
        p0.font.bold = True
        p0.space_after = Pt(6)

    for i, b in enumerate(bullets):
        p = tf.add_paragraph() if (i > 0 or (title_text is not None)) else tf.paragraphs[0]
        if i == 0 and title_text is not None:
            # ensure first bullet starts after title
            p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(font_pt)
        p.space_after = Pt(6)
        p.line_spacing = line_spacing
        p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT
