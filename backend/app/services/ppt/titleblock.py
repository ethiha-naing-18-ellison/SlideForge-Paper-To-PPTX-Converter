"""Title block creation utilities for clean, professional slide titles."""

from __future__ import annotations
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .cleantext import normalize_text, is_generic_title

def draw_title_block(slide, title: str, subtitle: str | None, authors: str | None, layout):
    """Draw a clean title block with title, subtitle, and authors."""
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    title_top = layout.margin_top_in
    title_h = layout.title_height_in

    # Title (big, wrapped)
    tb = slide.shapes.add_textbox(Inches(left), Inches(title_top), Inches(width), Inches(title_h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = normalize_text(title or "Untitled")
    p.font.size = Pt(34)
    p.font.bold = True

    # Subtitle (doc type + venue/year)
    if subtitle:
        sb = slide.shapes.add_textbox(
            Inches(left), 
            Inches(title_top + title_h + 0.1), 
            Inches(width), 
            Inches(0.55)
        )
        st = sb.text_frame
        st.word_wrap = True
        sp = st.paragraphs[0]
        sp.text = normalize_text(subtitle)
        sp.font.size = Pt(18)

    # Authors (small, gray footer line near top area)
    if authors:
        ab = slide.shapes.add_textbox(
            Inches(left), 
            Inches(title_top + title_h + 0.75), 
            Inches(width), 
            Inches(0.45)
        )
        at = ab.text_frame
        at.word_wrap = True
        ap = at.paragraphs[0]
        ap.text = normalize_text(authors)
        ap.font.size = Pt(12)
        ap.font.color.rgb = RGBColor(180, 186, 197)

def create_clean_title_slide(prs, layout, doc_title: str, doc_type_label: str, authors: str | None = None):
    """Create a clean title slide with proper formatting."""
    from .cleanup import remove_unused_placeholders
    
    # Use blank layout to avoid default placeholders
    try:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    except:
        slide = prs.slides.add_slide(prs.slide_layouts[0])  # Fallback to title layout
    
    # Clean up any default placeholders
    remove_unused_placeholders(slide)
    
    # Draw our custom title block
    draw_title_block(slide, doc_title, doc_type_label, authors, layout)
    
    return slide
