"""Callout utilities for displaying key terms and important information."""

from __future__ import annotations
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def clean_terms(terms: list[str]) -> list[str]:
    """Clean key terms by removing prefixes and normalizing."""
    out = []
    for term in terms or []:
        term = term.strip()
        # Remove common prefixes
        term = term.replace("keywords", "").replace("keyword", "").strip(": -")
        # Title case
        term = term.title()
        if term:
            out.append(term)
    return out[:5]  # Limit to 5 terms

def add_keyterms_callout(slide, layout, terms: list[str]):
    """Add a key terms callout box to the bottom-right of a slide."""
    # Clean the terms first
    terms = clean_terms(terms)
    if not terms:
        return
    
    # Limit to 5 terms max
    display_terms = terms[:5]
    text = "Key terms: " + ", ".join(display_terms)
    
    # Calculate position (bottom-right)
    w, h = 4.8, 0.7
    x = layout.slide_width_in - layout.margin_right_in - w
    y = layout.slide_height_in - layout.margin_bottom_in - h - 0.1
    
    # Create textbox
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(180, 220, 255)
    
    # Add subtle background and border
    try:
        fill = box.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(25, 35, 55)
        
        line = box.line
        line.color.rgb = RGBColor(60, 90, 140)
    except Exception:
        pass  # Skip styling if not supported

def add_highlight_box(slide, layout, text: str, position: str = "top-right"):
    """Add a highlight box with custom text at specified position."""
    if not text:
        return
    
    # Position calculations
    w, h = 3.5, 0.8
    if position == "top-right":
        x = layout.slide_width_in - layout.margin_right_in - w
        y = layout.margin_top_in
    elif position == "bottom-left":
        x = layout.margin_left_in
        y = layout.slide_height_in - layout.margin_bottom_in - h
    else:  # center
        x = (layout.slide_width_in - w) / 2
        y = layout.margin_top_in
    
    # Create textbox
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Add background
    try:
        fill = box.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 100, 200)
    except Exception:
        pass
